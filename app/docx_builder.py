"""
Converts adapted resume Markdown to a formatted .docx file.

Expected Markdown structure (matches Andrey's resume format):
  Line 1:  **NAME**                        → large bold name
  Line 2:  **Title | Tagline**             → subtitle
  Line 3:  contact links line              → contact (10pt)
  Line 4:  Location | *note*              → location
  blank
  **SECTION HEADER**                       → section (bold, separator line)
  paragraph text                           → body text
  **Company**   Location                   → company header
  **Date**                                 → date
  ***Role***                               → job title (bold+italic)
  *Description*                            → description (italic)
    - bullet                               → bullet point
  **Label:**  text                         → skill line (label bold + text)
"""

import io
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ── Colour palette ────────────────────────────────────────────────────────────
COLOR_BLACK   = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_DARK    = RGBColor(0x22, 0x22, 0x22)
COLOR_MID     = RGBColor(0x44, 0x44, 0x44)
COLOR_RULE    = RGBColor(0xAA, 0xAA, 0xAA)

FONT_NAME = "Calibri"


# ── Low-level helpers ──────────────────────────────────────────────────────────

def _set_font(run, size_pt: float, bold=False, italic=False,
              color: RGBColor = COLOR_DARK):
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _para_space(para, before_pt=0, after_pt=0, line_rule=None):
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(int(before_pt * 20)))
    spacing.set(qn("w:after"), str(int(after_pt * 20)))
    if line_rule:
        spacing.set(qn("w:lineRule"), line_rule[0])
        spacing.set(qn("w:line"), str(line_rule[1]))
    pPr.append(spacing)


def _add_bottom_border(para, color="AAAAAA", size=4):
    """Thin bottom rule below a paragraph (used for section headers)."""
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _strip_md(text: str) -> str:
    """Remove markdown bold/italic markers from a string."""
    return re.sub(r'\*{1,3}(.*?)\*{1,3}', r'\1', text)


def _strip_links(text: str) -> str:
    """Replace [label](url) with label."""
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)


def _clean(text: str) -> str:
    return _strip_links(_strip_md(text)).strip()


# ── Inline formatting (bold / italic / bold+italic) ───────────────────────────

_INLINE_RE = re.compile(r'(\*{1,3})(.*?)\1|(\[([^\]]+)\]\([^)]+\))|(<([^>]+)>)')


def _add_inline(paragraph, text: str, base_size: float,
                base_bold=False, base_italic=False,
                base_color: RGBColor = COLOR_DARK):
    """Parse inline ** / * / *** markers and add runs with proper formatting."""
    pos = 0
    for m in _INLINE_RE.finditer(text):
        # plain text before this match
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            _set_font(run, base_size, base_bold, base_italic, base_color)

        stars = m.group(1)
        inner = m.group(2)
        link_full = m.group(3)
        link_text = m.group(4)
        email_full = m.group(5)
        email_text = m.group(6)

        if stars is not None:
            bold   = base_bold   or len(stars) >= 2
            italic = base_italic or len(stars) in (1, 3)
            run = paragraph.add_run(inner)
            _set_font(run, base_size, bold, italic, base_color)
        elif link_text:
            run = paragraph.add_run(link_text)
            _set_font(run, base_size, base_bold, base_italic, base_color)
        elif email_text:
            run = paragraph.add_run(email_text)
            _set_font(run, base_size, base_bold, base_italic, base_color)

        pos = m.end()

    # remaining plain text
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_font(run, base_size, base_bold, base_italic, base_color)


# ── Line classifiers ──────────────────────────────────────────────────────────

SECTION_HEADERS = {
    "PROFILE", "PROFESSIONAL EXPERIENCE", "CORE SKILLS",
    "CERTIFICATIONS & EDUCATION", "CERTIFICATIONS", "EDUCATION",
    "SKILLS", "EXPERIENCE", "SUMMARY",
}

def _is_section_header(line: str) -> bool:
    cleaned = re.sub(r'\*+', '', line).strip().upper()
    if cleaned in SECTION_HEADERS:
        return True
    # all-caps bold line
    if re.match(r'^\*\*[A-Z][A-Z\s&/]+\*\*$', line.strip()):
        return True
    return False


def _is_name_line(line: str) -> bool:
    return bool(re.match(r'^\*\*[A-Z][A-Z\s]+\*\*\s*$', line.strip()))


def _is_subtitle(line: str) -> bool:
    return bool(re.match(r'^\*\*[^*]+\*\*\s*$', line.strip())) and '|' in line


def _is_role_line(line: str) -> bool:
    return bool(re.match(r'^\*\*\*[^*]+\*\*\*\s*$', line.strip()))


def _is_italic_desc(line: str) -> bool:
    return bool(re.match(r'^\*[^*].*\*\s*$', line.strip()))


def _is_bullet(line: str) -> bool:
    return bool(re.match(r'^\s*[-•]\s+', line))


def _is_skill_line(line: str) -> bool:
    return bool(re.match(r'^\*\*[^*]+:\*\*\s+', line.strip()))


def _is_company_line(line: str) -> bool:
    """Bold company name optionally followed by location."""
    return bool(re.match(r'^\*\*[^*]+\*\*\s{2,}', line.strip()))


def _is_date_line(line: str) -> bool:
    return bool(re.match(r'^\*\*(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|\d{4})', line.strip()))


# ── Main builder ──────────────────────────────────────────────────────────────

def build_docx(markdown_text: str) -> bytes:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin   = Cm(2.0)
        section.right_margin  = Cm(2.0)

    # Remove default empty paragraph
    for p in doc.paragraphs:
        p._element.getparent().remove(p._element)

    lines = markdown_text.strip().split('\n')
    first_nonblank = True
    line_index = 0

    while line_index < len(lines):
        raw = lines[line_index]
        line = raw.strip()
        line_index += 1

        if not line:
            continue

        # ── Name (first bold-only all-caps line) ──────────────────────────────
        if first_nonblank and _is_name_line(line):
            first_nonblank = False
            para = doc.add_paragraph()
            _para_space(para, before_pt=0, after_pt=2)
            name_text = _clean(line)
            run = para.add_run(name_text.upper())
            _set_font(run, 20, bold=True, color=COLOR_BLACK)
            continue

        first_nonblank = False

        # ── Subtitle / Title ──────────────────────────────────────────────────
        if _is_subtitle(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=0, after_pt=1)
            _add_inline(para, line, 11, base_color=COLOR_MID)
            continue

        # ── Contact line (contains • or @ or linkedin) ────────────────────────
        if any(x in line for x in ['•', '@', 'linkedin', 'github']) and not _is_section_header(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=0, after_pt=1)
            _add_inline(para, line, 10, base_color=COLOR_MID)
            continue

        # ── Location line ─────────────────────────────────────────────────────
        if line.startswith('Berlin') or line.startswith('London') or (
            'Germany' in line or 'Authorized' in line
        ) and not _is_section_header(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=0, after_pt=6)
            _add_inline(para, line, 10, base_color=COLOR_MID)
            continue

        # ── Section header ────────────────────────────────────────────────────
        if _is_section_header(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=10, after_pt=3)
            header_text = _clean(line).upper()
            run = para.add_run(header_text)
            _set_font(run, 10.5, bold=True, color=COLOR_BLACK)
            _add_bottom_border(para, color="BBBBBB", size=4)
            continue

        # ── Job role (***bold italic***) ──────────────────────────────────────
        if _is_role_line(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=2, after_pt=1)
            role_text = _clean(line)
            run = para.add_run(role_text)
            _set_font(run, 11, bold=True, italic=True, color=COLOR_DARK)
            continue

        # ── Company line (**Company**   Location) ─────────────────────────────
        if _is_company_line(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=8, after_pt=0)
            _add_inline(para, line, 11, base_color=COLOR_DARK)
            continue

        # ── Date line (**Month YYYY – ...**) ─────────────────────────────────
        if _is_date_line(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=0, after_pt=1)
            date_text = _clean(line)
            run = para.add_run(date_text)
            _set_font(run, 10, bold=False, color=COLOR_MID)
            continue

        # ── Italic description line ───────────────────────────────────────────
        if _is_italic_desc(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=1, after_pt=2)
            _add_inline(para, line, 10, base_italic=True, base_color=COLOR_MID)
            continue

        # ── Skill line (**Label:**  text) ─────────────────────────────────────
        if _is_skill_line(line):
            para = doc.add_paragraph()
            _para_space(para, before_pt=2, after_pt=1)
            # Split label from rest
            m = re.match(r'^(\*\*[^*]+:\*\*)\s+(.*)', line.strip())
            if m:
                label_md, rest = m.group(1), m.group(2)
                label_text = _clean(label_md)
                run = para.add_run(label_text + "  ")
                _set_font(run, 10, bold=True, color=COLOR_DARK)
                run2 = para.add_run(rest)
                _set_font(run2, 10, color=COLOR_DARK)
            else:
                _add_inline(para, line, 10, base_color=COLOR_DARK)
            continue

        # ── Bullet point ──────────────────────────────────────────────────────
        if _is_bullet(line):
            para = doc.add_paragraph(style="List Bullet")
            _para_space(para, before_pt=1, after_pt=1)
            # Set indent
            pPr = para._p.get_or_add_pPr()
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "360")
            ind.set(qn("w:hanging"), "180")
            pPr.append(ind)

            bullet_text = re.sub(r'^\s*[-•]\s+', '', line)
            _add_inline(para, bullet_text, 10, base_color=COLOR_DARK)

            # Check for continuation lines (indented, no bullet)
            while line_index < len(lines):
                next_raw = lines[line_index]
                next_stripped = next_raw.strip()
                if not next_stripped:
                    break
                # continuation: indented and not a new bullet / section / etc.
                if (next_raw.startswith('    ') or next_raw.startswith('\t')) and not _is_bullet(next_stripped):
                    run = para.add_run(' ' + next_stripped)
                    _set_font(run, 10, color=COLOR_DARK)
                    line_index += 1
                else:
                    break
            continue

        # ── Default: body paragraph ───────────────────────────────────────────
        para = doc.add_paragraph()
        _para_space(para, before_pt=2, after_pt=2)
        _add_inline(para, line, 10.5, base_color=COLOR_DARK)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
