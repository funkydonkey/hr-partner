from app.docx_builder import build_docx
from docx import Document
import io


SAMPLE_MD = """\
**ANDREY MOLCHANSKY**

**OneStream Lead Architect | OCP Certified | CPM Platform & Team Leadership**

website • +49 176 123 • email@example.com • linkedin.com/in/test

Berlin, Germany | *Authorized to work in Germany*

**PROFILE**

Experienced OneStream architect with 5 years leading EPM transformations.

**PROFESSIONAL EXPERIENCE**

**Acme Corp**   Berlin, Germany
**Jan 2022 – Present**

***OneStream Lead Architect***

*Global company; OneStream runs Group consolidation for 100 entities.*

  - Led migration from legacy system to OneStream.
  - Designed dashboards and workflows.

**CORE SKILLS**

**CPM:**  OneStream, Anaplan, Power BI

**CERTIFICATIONS & EDUCATION**

  - OneStream Certified Professional (2024).
"""


def test_build_docx_returns_bytes():
    result = build_docx(SAMPLE_MD)
    assert isinstance(result, bytes)
    assert len(result) > 5000


def test_docx_is_valid_document():
    result = build_docx(SAMPLE_MD)
    doc = Document(io.BytesIO(result))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Name should appear
    assert any("ANDREY MOLCHANSKY" in t.upper() for t in texts)
    # Section headers
    assert any("PROFILE" in t for t in texts)
    assert any("PROFESSIONAL EXPERIENCE" in t for t in texts)


def test_docx_has_correct_structure():
    result = build_docx(SAMPLE_MD)
    doc = Document(io.BytesIO(result))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Should have more than 5 paragraphs
    assert len(texts) > 5
